from __future__ import annotations

import io
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def merge_pdfs(sources: list[Path], output: Path) -> None:
    merged = fitz.open()
    try:
        for source_path in sources:
            try:
                with fitz.open(source_path) as source:
                    merged.insert_pdf(source)
            except Exception as exc:
                raise ValueError(f"Could not read {source_path.name} as a PDF.") from exc
        merged.save(output, garbage=3, deflate=True)
    finally:
        merged.close()


def _set_section_size(section, width_pt: float, height_pt: float) -> None:
    section.page_width = Inches(width_pt / 72)
    section.page_height = Inches(height_pt / 72)
    section.top_margin = Inches(0)
    section.bottom_margin = Inches(0)
    section.left_margin = Inches(0)
    section.right_margin = Inches(0)


def _page_png(page) -> bytes:
    return page.get_pixmap(matrix=fitz.Matrix(180 / 72, 180 / 72), alpha=False).tobytes("png")


def _text_page(document: Document, page) -> None:
    for line in page.get_text("text").splitlines():
        line = line.strip()
        if line:
            paragraph = document.add_paragraph(line)
            paragraph.paragraph_format.space_after = Pt(2)


def convert_pdf_to_docx(source: Path, output: Path, mode: str) -> None:
    with fitz.open(source) as pdf:
        if pdf.is_encrypted:
            raise ValueError("This PDF is password protected. Remove its password before converting it.")
        document = Document()
        if document.paragraphs:
            node = document.paragraphs[0]._element
            node.getparent().remove(node)
        for index, page in enumerate(pdf):
            section = document.sections[0] if index == 0 else document.add_section(WD_SECTION.NEW_PAGE)
            width, height = page.rect.width, page.rect.height
            _set_section_size(section, width, height)
            if mode == "editable":
                _text_page(document, page)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.add_run().add_picture(io.BytesIO(_page_png(page)), width=Inches(width / 72), height=Inches(height / 72))
                if mode == "hybrid":
                    text = page.get_text("text").strip()
                    if text:
                        document.add_page_break()
                        _text_page(document, page)
    document.save(output)
